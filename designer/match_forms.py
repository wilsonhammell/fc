"""Which row of the business's forms workbook each final form is.

    python designer/match_forms.py

The designer's "Import Configuration from Form Table" fills a form's Config
and its signing buttons from the row a person picks by name. This picks it
the same way, from designer/form_table.json (the rows as the designer reads
them, written by form_table.mjs), in this order:

    1. the form's production name or file name, exactly, on the sheet the
       designer opens on (Confirmed Forms)
    2. the same on any other sheet
    3. a HISTORICAL FORM NAME(S) entry, exactly (read here, since the
       designer does not read that column)
    4. a candidate: the closest name, once common shortenings are spelt out
       ("Ln" -> "loan", "Agremnt" -> "agreement"), at 0.86 or better

A form matching none of them is listed as unmatched; make_handover.py puts
it in a folder of its own. Writes designer/form_matches.json and
designer/form_matches.csv (for a person to check the candidates).
"""
import csv
import difflib
import json
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK = r"C:\Users\Wilson\Desktop\form_requirements\Combined Form Requirements.xlsx"
TABLE = os.path.join(HERE, "designer", "form_table.json")
PILOT = os.path.join(HERE, "pilot_forms.txt")
NAMING = os.path.join(HERE, "field naming")

SPELT = {"ln": "loan", "agremnt": "agreement", "agrement": "agreement", "ben": "beneficiary",
         "sh": "successor holder", "info": "information", "req": "request", "acct": "account",
         "ins": "insurance", "trans": "transfer", "reimbursment": "reimbursement",
         "equip": "equipment", "atm": "atm", "abm": "atm"}
CANDIDATE = 0.86
# rows that are a decision about a form rather than a form: the workbook puts
# "Not Required" or "To be removed from CRM" in NEW FORM NAME and keeps the
# form's own name under HISTORICAL FORM NAME(S)
STATUS = ("not required", "to be removed from crm")


def norm(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def spelt(text):
    return " ".join(SPELT.get(w, w) for w in norm(text).split())


def pilot():
    """[(production name, file stem)] for the forms that have a file."""
    out = []
    for line in open(PILOT, encoding="utf-8"):
        line = line.strip()
        if not line or line.startswith("#") or "=>" not in line:
            continue
        name, stem = (p.strip() for p in line.split("=>", 1))
        if stem:
            out.append((name, stem))
    return out


def historical():
    """{(sheet, row): [historical names]} for every sheet that has the column."""
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True, max_row=ws.max_row))
        for r, row in enumerate(rows[:10], 1):
            heads = [norm(h) for h in row]
            col = next((i for i, h in enumerate(heads) if h.startswith("histor")), None)
            if col is None:
                continue
            for n, data in enumerate(rows[r:], r + 1):
                if data and col < len(data) and data[col]:
                    out[(ws.title, n)] = [p.strip() for p in re.split(r"[\n;]+", str(data[col]))
                                          if p.strip()]
            break
    return out


def main():
    table = json.load(open(TABLE, encoding="utf-8"))
    first = table["defaultSheet"]
    status_rows = [r for r in table["rows"] if norm(r["formName"]) in STATUS]
    rows = [r for r in table["rows"] if norm(r["formName"]) not in STATUS]
    hist = historical()
    named_folders = {d for d in os.listdir(NAMING) if os.path.isfile(os.path.join(NAMING, d, "fields.csv"))}
    results = []
    for prod, stem in pilot():
        keys = {norm(prod), norm(stem)}
        hit, how, score = None, "", 0.0
        on_first = [r for r in rows if r["sheet"] == first and norm(r["formName"]) in keys]
        others = [r for r in rows if r["sheet"] != first and norm(r["formName"]) in keys]
        by_hist = [r for r in rows if any(norm(h.removesuffix(".pdf")) in keys
                                          for h in hist.get((r["sheet"], r["row"]), []))]
        if on_first:
            hit, how, score = on_first[0], "exact", 1.0
        elif others:
            hit, how, score = others[0], "exact, other sheet", 1.0
        elif by_hist:
            hit = next((r for r in by_hist if r["sheet"] == first), by_hist[0])
            how, score = "historical name", 1.0
        else:
            want = [spelt(prod), spelt(stem)]
            best = max(((max(difflib.SequenceMatcher(None, w, spelt(r["formName"])).ratio()
                             for w in want), r) for r in rows),
                       key=lambda t: (t[0], t[1]["sheet"] == first))
            hit_name = ""
            if best[0] >= CANDIDATE:
                score, hit, how = best[0], best[1], "candidate"
            else:
                score, how = best[0], "unmatched"
                hit_name = best[1]["formName"]
        # a form the workbook has decided against, named under its old name
        decided = next((r for r in status_rows if any(norm(h.removesuffix(".pdf")) in keys
                        for h in hist.get((r["sheet"], r["row"]), []))), None)
        if how == "unmatched" and decided:
            how = f"unmatched - the workbook says {decided['formName']!r}"
            hit_name = f"{decided['sheet']} row {decided['row']}"
        folder = next((d for d in named_folders if norm(d) == norm(stem)), None)
        results.append({
            "production name": prod, "file": stem, "folder": folder, "match": how,
            "score": round(score, 3),
            "sheet": hit["sheet"] if hit else "", "row": hit["row"] if hit else "",
            "row name": hit["formName"] if hit else "",
            "closest": "" if hit else hit_name,
            "config": hit["config"] if hit else {}, "buttons": hit["buttons"] if hit else {},
            "notes": hit["notes"] if hit else [],
        })
    out = os.path.join(HERE, "designer", "form_matches.json")
    json.dump(results, open(out, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    with open(os.path.join(HERE, "designer", "form_matches.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        wr = csv.writer(fh)
        wr.writerow(["form", "file", "match", "score", "sheet", "row", "row in the workbook", "closest name"])
        for r in results:
            wr.writerow([r["production name"], r["file"], r["match"], r["score"], r["sheet"],
                         r["row"], r["row name"], r["closest"]])
    from collections import Counter
    print(dict(Counter(r["match"] for r in results)), "of", len(results))
    for r in results:
        if r["match"] != "exact":
            print(f"   {r['match']:20} {r['production name'][:48]:48} -> "
                  f"{(r['row name'] or r['closest'])[:50]} ({r['sheet'] or '-'} {r['row']}, {r['score']})")
        if not r["folder"]:
            print(f"   NO FIELD NAMING FOLDER for {r['file']}")


if __name__ == "__main__":
    main()
